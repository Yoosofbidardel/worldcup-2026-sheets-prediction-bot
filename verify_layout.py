"""Offline sanity check of the sheet-layout assumptions in config.py.

Runs against the local .xlsx with openpyxl (no Google credentials needed) so you
can confirm the slot/match math before connecting the live bot.

    python3 verify_layout.py "../مکا worldcup 2026.xlsx"
"""
import sys

import openpyxl
from openpyxl.utils import get_column_letter as L

from config import (
    EXCLUDED_NAMES,
    FIRST_SLOT_COL,
    MATCH_FIRST_ROW,
    MATCH_LAST_ROW,
    NAME_ROW,
    SLOT_STRIDE,
    SPECIAL_ROWS,
    WORKSHEET_NAME,
)


def blank(v):
    return v is None or str(v).strip() == ""


def main(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb[WORKSHEET_NAME] if WORKSHEET_NAME in wb.sheetnames else wb.worksheets[0]

    print(f"Worksheet: {ws.title}\n")
    print("PARTICIPANT SLOTS (idx — name — home/away cols — formula col present?)")
    col, idx = FIRST_SLOT_COL, 1
    while col <= ws.max_column:
        name = ws.cell(NAME_ROW, col).value
        if name and str(name).strip() and str(name).strip() not in EXCLUDED_NAMES:
            f = ws.cell(MATCH_FIRST_ROW, col + 2).value
            has_formula = isinstance(f, str) and f.startswith("=")
            print(f"  {idx:2d}. {str(name).strip():14s} "
                  f"{L(col)}/{L(col+1)}  formula@{L(col+2)}={'OK' if has_formula else 'MISSING'}")
            idx += 1
        col += SLOT_STRIDE
    print(f"  -> {idx-1} assignable slots\n")

    print("MATCHES (with both teams set)")
    n_open = n_done = 0
    for r in range(MATCH_FIRST_ROW, MATCH_LAST_ROW + 1):
        home, away = ws.cell(r, 1).value, ws.cell(r, 2).value
        if blank(home) or blank(away):
            continue
        ah, aw = ws.cell(r, 3).value, ws.cell(r, 4).value
        done = not blank(ah) and not blank(aw)
        n_done += done
        n_open += not done
        flag = f"{ah}-{aw}" if done else "open"
        print(f"  row {r}: {str(home).strip()} – {str(away).strip()}  [{flag}]")
    print(f"  -> {n_open} open, {n_done} finished\n")

    print("SPECIALS")
    for row, (label, pts) in SPECIAL_ROWS.items():
        actual = ws.cell(row, 3).value
        f = ws.cell(row, FIRST_SLOT_COL + 2).value  # formula in first slot's points col
        has_formula = isinstance(f, str) and f.startswith("=")
        state = "open" if blank(actual) else f"answer={str(actual).strip()}"
        print(f"  row {row}: {label} ({pts} pts) [{state}] formula={'OK' if has_formula else 'n/a'}")


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else "../مکا worldcup 2026.xlsx"
    main(path)
